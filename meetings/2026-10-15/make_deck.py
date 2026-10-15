"""Progress-report deck, 2026-10-15. Numbers: review-ai/docs/thesis_numbers.md (§1★, §4–§6d) and review-ai/results/.
Run:  PYTHONPATH=<dir with python-pptx> review-ai/.venv/bin/python make_deck.py   (then make_charts.py first)"""
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

D = Path(__file__).resolve().parent
RES = D.parent.parent / "review-ai" / "results"
FONT = "Yu Gothic"
INK, INK2, MUTED = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x8A, 0x89, 0x83)
ACCENT, ACCENT_BG = RGBColor(0x1C, 0x5C, 0xAB), RGBColor(0xEA, 0xF2, 0xFC)
WARN, WARN_BG = RGBColor(0xB5, 0x47, 0x1C), RGBColor(0xFD, 0xF0, 0xE8)
ROW_ALT, WHITE = RGBColor(0xF5, 0xF4, 0xF1), RGBColor(0xFF, 0xFF, 0xFF)


def load(rel):
    p = RES / rel
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


INCLUDE_XL = True  # result finished 2026-10-14 22:45; set False to report it as still running
XL = load("followup/B_xlmr_large_E7E8_vs_E1.json") if INCLUDE_XL else None
if XL:
    _m = XL["metrics"]["agg_neg"]
    _v = [json.loads((RES / f"{p}_metrics.json").read_text(encoding="utf-8"))["valid_agg_negative_f1"]
          for p in [f"CMP_E1_xlmr_large_seed{s}" for s in (42, 43, 44)] + [f"B_E8_xlmr_large_seed{s}" for s in (42, 43, 44)]]
    XL_DELTA = f"{100 * _m['delta_point']:+.2f}" + ("" if _m["significant_at_95"] else "（有意差なし）")
    XL_VALID = "逆" if sum(_v[3:]) < sum(_v[:3]) else "同じ"
    XL_SCORE = f"{_m['b_mean']:.3f}"
    assert not _m["significant_at_95"], "slide 7 wording assumes no significant change on XLM-R large"
else:
    XL_DELTA, XL_VALID, XL_SCORE = "実行中（結果は後日ご報告）", "", "実行中"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def style(run, size=16, bold=False, color=INK):
    f = run.font
    f.size, f.bold, f.color.rgb, f.name = Pt(size), bold, color, FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text(slide, x, y, w, h, paras, size=16, color=INK, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, spacing=6):
    """paras: list of str | (str, dict) where dict may hold size/bold/color/level/after."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, p in enumerate(paras):
        s, o = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(o.get("after", spacing))
        lvl = o.get("level", 0)
        prefix = {0: "", 1: "・", 2: "　– "}[lvl] if o.get("bullet", lvl > 0) else ""
        r = para.add_run()
        r.text = prefix + s
        style(r, o.get("size", size - 2 * (lvl == 2)), o.get("bold", False), o.get("color", color))
    return tb


def box(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    return shp


def header(slide, title, num, kicker=None):
    box(slide, 0, 0, 0.18, 7.5, ACCENT)
    if kicker:
        text(slide, 0.6, 0.28, 11, 0.4, [(kicker, {"size": 13, "color": ACCENT, "bold": True})])
    text(slide, 0.6, 0.55, 12, 0.8, [(title, {"size": 28, "bold": True})])
    text(slide, 12.3, 7.0, 0.8, 0.35, [(str(num), {"size": 11, "color": MUTED})], align=PP_ALIGN.RIGHT)


def table(slide, x, y, w, rows, col_w, size=14, row_h=0.46, highlight=(), sub_small=True):
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))).table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = ACCENT
            elif i in highlight:
                c.fill.fore_color.rgb = ACCENT_BG
            else:
                c.fill.fore_color.rgb = ROW_ALT if i % 2 == 0 else WHITE
            tf = c.text_frame
            tf.word_wrap = True
            for k, ln in enumerate(val.split("\n")):
                para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                r = para.add_run()
                r.text = ln
                small = sub_small and k > 0
                style(r, size - 2 if small else size, bold=(i == 0 or (i in highlight and j == 0)),
                      color=WHITE if i == 0 else (INK2 if small else INK))
    return t


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def callout(slide, x, y, w, h, title, body, warn=False, size=14):
    fg, bg = (WARN, WARN_BG) if warn else (ACCENT, ACCENT_BG)
    box(slide, x, y, w, h, bg)
    box(slide, x, y, 0.08, h, fg)
    text(slide, x + 0.25, y + 0.15, w - 0.4, h - 0.3,
         [(title, {"size": size + 1, "bold": True, "color": fg})] + [(b, {"size": size}) for b in body], spacing=4)


def foot(slide, s):
    text(slide, 0.6, 6.85, 11.5, 0.4, [(s, {"size": 11, "color": MUTED})])


# ---------------------------------------------------------------- 1 title
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, ACCENT)
text(s, 0.9, 1.3, 11.5, 0.5, [("研究進捗報告", {"size": 18, "color": RGBColor(0xCD, 0xE2, 0xFB), "bold": True})])
text(s, 0.9, 2.0, 11.6, 2.4, [
    ("評価点を弱教師とした事前学習による", {"size": 34, "bold": True, "color": WHITE, "after": 2}),
    ("日本語ホテルレビューのアスペクト感情分析", {"size": 34, "bold": True, "color": WHITE, "after": 10}),
    ("― 研究の現状と論文化の見通し ―（題目は仮）", {"size": 20, "color": RGBColor(0xCD, 0xE2, 0xFB)}),
])
text(s, 0.9, 5.6, 11, 1.0, [("2026年10月15日", {"size": 16, "color": WHITE, "after": 2}),
                            ("GU JINCHENG", {"size": 16, "color": WHITE})])
notes(s, "开场：今天报告研究到了哪里、能不能写成论文、有哪些要商量。")

# ---------------------------------------------------------------- 2 agenda
s = prs.slides.add_slide(BLANK)
header(s, "本日の報告", 2)
callout(s, 0.6, 1.6, 6.0, 2.6, "結論", [
    "予定していた実験はすべて完了し、論文として書ける段階に来ています。",
    "実験はここで締めて、執筆に移りたいと考えています。"], size=17)
text(s, 7.2, 1.6, 5.6, 4.6, [
    ("本日の内容", {"size": 20, "bold": True, "color": ACCENT, "after": 12}),
    ("1. 研究の概要", {"after": 10}),
    ("2. 主な結果（4点）", {"after": 10}),
    ("3. 効く条件と限界", {"after": 10}),
    ("4. 論文の構成と計画", {"after": 10}),
    ("5. ご相談したいこと", {"after": 10}),
], size=18)
notes(s, "先说结论：实验已完成，可以写论文，想转入写作。")

# ---------------------------------------------------------------- 3 overview
s = prs.slides.add_slide(BLANK)
header(s, "研究の概要", 3)
cols = [
    ("課題", [
        "文ごとに「どの観点について、肯定か否定か」を判定する（7観点×正負＝14ラベル）",
        "ネガティブはラベル全体の26%しかない",
        "ネガティブの92%は、感情辞書の否定語を含まない（暗示的）",
        "例：「駅から歩いて20分かかった。」→ 立地・ネガティブ（説明用の作例）"]),
    ("アイデア", [
        "楽天トラベルのレビューには、投稿者が付けた観点別の評価点がある（6観点、1〜5点）",
        "約700万件。追加のアノテーション費用はかからない",
        "これを弱い教師信号として使う"]),
    ("手法と評価", [
        "① 評価点で事前学習：1〜2点をネガティブ、3〜5点をそれ以外、未評価は学習に使わない",
        "② 人手ラベルの Gold データ（53,192文）で微調整",
        "比較対象：①なしで②だけ（事前学習なし）",
        "評価：テスト5,356文のネガティブ集約F1。複数回の実行と paired bootstrap で95%信頼区間"]),
]
for i, (h, items) in enumerate(cols):
    x = 0.6 + i * 4.15
    box(s, x, 1.55, 3.95, 5.0, ROW_ALT)
    text(s, x + 0.2, 1.7, 3.6, 4.8, [(h, {"size": 18, "bold": True, "color": ACCENT, "after": 10})]
         + [(t, {"level": 1, "size": 14, "after": 8}) for t in items])
foot(s, "Gold：Nakayama et al.（LREC 2022）。評価点データ：楽天トラベル 2009〜2024年（国立情報学研究所 IDR 経由で提供）。Gold と重複するレビューは除外。")
notes(s, "背景（负面少且隐式）→ 想法（分项评分=免费弱监督）→ 两阶段方法与评价方式。例句是自己造的。")

# ---------------------------------------------------------------- 4 main result
s = prs.slides.add_slide(BLANK)
header(s, "結果① 評価点による事前学習はネガティブの検出を改善する", 4, kicker="主結果（XLM-R base）")
table(s, 0.6, 1.55, 12.1, [
    ["指標（テスト）", "事前学習なし（5回）", "事前学習あり（9回）", "差", "95% 信頼区間"],
    ["ネガティブ集約 F1", "0.7066", "0.7257", "+1.91", "[+1.10, +2.75]"],
    ["暗示的ネガティブ（2,065）", "0.8403", "0.8546", "+1.43", "[+0.69, +2.18]"],
    ["ポジティブ集約 F1", "0.8605", "0.8623", "+0.17", "有意差なし"],
], [3.6, 2.55, 2.55, 1.2, 2.2], size=15, row_h=0.55, highlight=(1,), sub_small=False)
text(s, 0.6, 3.85, 7.4, 3.0, [
    ("事前学習ありの9回すべてが、事前学習なしの最高値（0.7143）を上回る", {"level": 1, "after": 10}),
    ("観点別：立地 +6.07、風呂 +4.06、部屋 +3.56、設備 +2.42（有意）。サービス・朝食・夕食は有意差なし", {"level": 1, "after": 10}),
    ("適合率と再現率がともに改善し、ネガティブと予測する数はほぼ同じ（2,353 → 2,363）", {"level": 1, "after": 10}),
], size=15)
callout(s, 8.4, 3.85, 4.3, 2.4, "読み取れること", [
    "改善はネガティブに限られる。", "Gold で数が少ないクラスを、評価点が補っている。"])
notes(s, "主结果：负面 +1.91 显著，9 次全部高于基线最好值；隐式负面也显著；正面不变。精确率召回率都升，不是多报负面。")

# ---------------------------------------------------------------- 5 label efficiency
s = prs.slides.add_slide(BLANK)
header(s, "結果② 人手ラベルが半分でも、事前学習なし・全量と同程度", 5, kicker="新規：ラベル量を変えた実験")
s.shapes.add_picture(str(D / "chart_label_efficiency.png"), Inches(0.5), Inches(1.5), height=Inches(5.0))
text(s, 8.5, 1.6, 4.4, 3.4, [
    ("Gold の量を 10〜100% に変えても、改善幅は約 1.3〜1.9 ポイントでほぼ一定", {"level": 1, "after": 10}),
    ("事前学習ありなら Gold 50% でも、事前学習なし・全量と有意差なし（+0.47、信頼区間は 0 を含む）", {"level": 1, "after": 10}),
    ("内挿すると約 37% で全量の水準に届く（参考値）", {"level": 1, "after": 10}),
], size=15)
callout(s, 8.5, 5.0, 4.3, 1.55, "意味", ["評価点を使えば、人手ラベルを約半分に減らしても同じ水準を保てる。"])
foot(s, "各点 5 回の平均（ひげは標準偏差）。少量データで基線が学習不足にならないよう、更新回数を全量とそろえた設定（検証データで選択）。")
notes(s, "新实验：Gold 量 10–100%，增益恒定约 1.3–1.9；有预训练时一半 Gold 与无预训练全量无显著差。若固定 3 轮，10% 会显示 +5.0，但那是基线训练不足，按预注册规则不采用。")

# ---------------------------------------------------------------- 6 aspect-specific signal
s = prs.slides.add_slide(BLANK)
header(s, "結果③ 効いているのは「観点別」の評価点", 6, kicker="新規：信号の種類だけを変えた対照実験")
table(s, 0.6, 1.55, 7.3, [
    ["DeBERTa・2021〜24年のデータ", "観点別の評価点\n（6出力）", "総合評価点のみ\n（1出力）"],
    ["二値分類（BCE）", "0.7447\n事前学習なしより +1.85", "0.6818\n事前学習なしより −4.44"],
    ["対照学習（SCL）", "未実施", "0.7259\n事前学習なしより −0.03"],
], [2.9, 2.2, 2.2], size=15, row_h=0.95)
text(s, 0.6, 4.6, 7.3, 0.5, [("事前学習なし：0.7262（テストのネガティブ集約 F1、各3回）", {"size": 12, "color": INK2})])
text(s, 8.3, 1.55, 4.5, 3.5, [
    ("同じ学習方法のまま、信号を総合評価点に替えると、改善が消えるだけでなく悪化する（−4.44、有意）", {"level": 1, "after": 10}),
    ("観点別と総合の差：+6.29（DeBERTa）、+3.82（XLM-R）。DeBERTa では 7 観点すべてで観点別が上", {"level": 1, "after": 10}),
    ("総合評価点だけなら対照学習の方が良い（+4.41）。「二値分類が優れている」とは言えない", {"level": 1, "after": 10}),
], size=14)
callout(s, 0.6, 5.2, 12.1, 1.35, "読み取れること", [
    "効いているのは「どの観点が悪かったか」という情報。総合評価点のマルチタスク学習に効果がなかった先行研究（小泉 2025）とも整合する。"])
notes(s, "新对照实验：只换信号。分项 +1.85，总评分反而 −4.44；分项比总评分高 6.29（DeBERTa）/3.82（XLM-R）。但总评分下 SCL 比 BCE 好，所以不能说二分类优于对比学习。")

# ---------------------------------------------------------------- 7 backbones
s = prs.slides.add_slide(BLANK)
header(s, "結果④ 強いモデルほど改善は小さく、最も強いモデルでは消える", 7, kicker="新規：6種類のモデルで検証")
s.shapes.add_picture(str(D / "chart_backbones.png"), Inches(0.5), Inches(1.5), height=Inches(5.0))
text(s, 8.5, 1.6, 4.4, 3.8, [
    ("比較的弱いモデル（XLM-R base、DeBERTa）では明確に改善（約 +1.9）", {"level": 1, "after": 10}),
    ("BERT large と ModernBERT はテストでは有意だが、検証データでは逆方向", {"level": 1, "after": 10}),
    (f"XLM-R large と、最も強い LUKE-ja では効果なし（{XL_DELTA.replace('（有意差なし）', '')}、+0.15、いずれも有意差なし）"
     if XL else f"最も強い LUKE-ja では効果なし（+0.15、有意差なし）。XLM-R large：{XL_DELTA}", {"level": 1, "after": 10}),
    (f"同じ XLM-R でも、base は +1.91、large は {XL_DELTA.replace('（有意差なし）', '')}" if XL else "", {"level": 1, "after": 10}),
], size=15)
callout(s, 8.5, 5.25, 4.3, 1.3, "論文での書き方", ["「どのモデルでも効く」とは書かない。強いモデルでの限界として明記する。"], warn=True)
foot(s, "事前学習なしの値が大きいほど強いモデル。XLM-R base と DeBERTa は主結果・既存の比較、その他は 3 回ずつ。")
notes(s, "新实验：换 6 种底座。弱的明确有效，越强越小，最强的 LUKE 无效。要作为局限明确写。")

# ---------------------------------------------------------------- 8 comparison
s = prs.slides.add_slide(BLANK)
header(s, "他のモデルとの比較", 8, kicker="同じデータ分割・同じ評価指標")
table(s, 0.6, 1.55, 12.1, [
    ["種類", "モデルとテストのネガティブ集約 F1"],
    ["微調整したエンコーダ\n（事前学習なし、各3回）",
     "mBERT 0.662 / XLM-R base 0.708 / DeBERTa-v3-ja 0.724 / BERT large 0.739\nModernBERT-ja 310M 0.746 / XLM-R large 0.749 / LUKE-ja 0.756"],
    ["＋評価点による事前学習",
     f"XLM-R base 0.726 / DeBERTa-v3-ja 0.745 / BERT large 0.750\nModernBERT-ja 310M 0.756 / LUKE-ja 0.757 / XLM-R large {XL_SCORE}"],
    ["ローカル LLM\n（学習なし、指示＋20例）", "llm-jp-3.1-13B 0.489 / Qwen2.5-7B 0.597 / gpt-oss-20B 0.613 / Swallow-8B 0.629"],
], [3.3, 8.8], size=15, row_h=0.95, sub_small=False)
text(s, 0.6, 5.5, 12.1, 1.3, [
    ("モデルの違い（約 9 ポイント）は、事前学習による改善（1〜2 ポイント）より大きい", {"level": 1, "after": 8}),
    ("LLM は学習なしでは、最も古典的な mBERT にも届かない。GPT-4o 等の外部サービスはデータ利用規約のため未使用（ご相談事項）", {"level": 1, "after": 8}),
], size=15)
foot(s, "事前学習なしのエンコーダは比較実験での再実行（3回）。主結果（5回）とは実行が異なる。LLM は各モデルの良い方の設定。")
notes(s, "换底座的影响（约 9 点）大于本手法（1–2 点）；本地 LLM 不训练时低于 mBERT；GPT-4o 因条款未用。")

# ---------------------------------------------------------------- 9 conditions
s = prs.slides.add_slide(BLANK)
header(s, "効く条件と、効かなかった試み", 9)
table(s, 0.6, 1.55, 12.1, [
    ["変えた条件", "結果（事前学習なしとの差、ポイント）", "分かったこと"],
    ["事前学習のデータ量：58万 / 200万 / 676万件", "+1.0 / +3.1 / +2.4（各1回）", "約200万件で頭打ち"],
    ["データの年代：2021〜24年 / 2009〜13年", "+3.1 / +1.6（各1回）", "新しいデータほど有効"],
    ["本文に観点の語がある低評価だけを残す", "−1.08（有意）", "ノイズ除去はかえって有害"],
    ["食事の評価点を朝食・夕食に分ける", "朝食 −0.71、夕食 −0.16（5回、有意差なし）", "分けても改善しない"],
    ["改善が起きた文の分析（3モデル）", "暗示的ネガティブ全般で改善\n純粋に暗示的な文も +1.7 前後", "特定の種類の文に集中してはいない"],
], [4.6, 4.3, 3.2], size=14, row_h=0.62)
foot(s, "データ量・年代は各1回の実行のため参考扱い。同じ設定でも実行ごとに最大約 0.7 ポイントの揺れがある。")
notes(s, "条件：约 200 万饱和、近年更好；过滤噪声有害；餐食拆分无效；增益普遍，不集中于某类句子。")

# ---------------------------------------------------------------- 10 limitations
s = prs.slides.add_slide(BLANK)
header(s, "限界", 10)
text(s, 0.6, 1.6, 12.1, 5.0, [
    ("1. 強いモデルでは効果が小さい、または無い（XLM-R large、LUKE-ja）", {"after": 14}),
    ("2. 暗示的・明示的の区別は感情辞書による自動判定で、人手による検証はしていない", {"after": 14}),
    ("3. データセットは 1 つ（楽天トラベルのホテルレビュー）", {"after": 14}),
    ("4. サービス・朝食・夕食は改善しない（食事の評価点は 1 項目しかない）", {"after": 14}),
    ("5. データ量・年代の実験は各 1 回のみ", {"after": 14}),
    ("6. LLM との比較はローカルモデル（7B〜20B、学習なし）に限られる", {"after": 14}),
], size=18)
notes(s, "局限 6 条，论文里全部写明。")

# ---------------------------------------------------------------- 11 thesis plan
s = prs.slides.add_slide(BLANK)
header(s, "論文にできるか ― 構成と計画", 11)
callout(s, 0.6, 1.55, 5.9, 3.0, "判断：論文にできると考えます", [
    "主張 3 点（効果・ラベル効率・観点別の信号）が、事前に決めた検定で支持された",
    "限界も実験で明らかになっている",
    "すべての数字は、保存した予測から再計算できる"], size=15)
text(s, 0.6, 4.8, 5.9, 2.0, [
    ("スケジュール案（英語で執筆）", {"size": 17, "bold": True, "color": ACCENT, "after": 8}),
    ("10月後半：図表の作成、3〜5章", {"level": 1, "size": 15, "after": 4}),
    ("11月末：中間稿　／　12月末：初稿", {"level": 1, "size": 15, "after": 4}),
    ("提出日に合わせて調整します", {"level": 1, "size": 15, "after": 4}),
])
table(s, 7.0, 1.55, 5.7, [
    ["章（英語）", "内容"],
    ["1 Introduction", "課題と貢献"],
    ["2 Related Work", "ABSA、弱教師、暗示的感情"],
    ["3 Data", "Gold と評価点データ、重複除去"],
    ["4 Method", "2 段階の学習"],
    ["5 Experimental Setup", "指標、複数回実行と検定"],
    ["6 Results", "主結果、ラベル効率、信号、モデル"],
    ["7 Analysis", "条件、改善の出どころ、LLM 比較"],
    ["8 Discussion", "解釈と限界"],
    ["9 Conclusion", "まとめと今後の課題"],
], [2.4, 3.3], size=13, row_h=0.5)
notes(s, "能写成论文：三个主张有预注册检验支持，局限也清楚，数字可复算。章节结构与日程。")

# ---------------------------------------------------------------- 12 consult
s = prs.slides.add_slide(BLANK)
header(s, "ご相談したいこと", 12)
text(s, 0.6, 1.55, 12.1, 5.2, [
    ("1. 実験をここで締めて、執筆に移ってよいでしょうか", {"bold": True, "after": 2}),
    ("残る候補：観点別評価点＋対照学習、データ量実験の回数追加、暗示・明示の人手検証", {"level": 2, "after": 12}),
    ("2. 「強いモデルでは効果がない」ことを、論文でどの程度前面に出すべきでしょうか", {"bold": True, "after": 12}),
    ("3. GPT-4o 等の外部 API との比較：規約上、楽天の許可が必要です。申請すべきでしょうか", {"bold": True, "after": 12}),
    ("4. データ利用規約について", {"bold": True, "after": 2}),
    ("論文・発表の前の楽天への通知、研究終了時のデータ削除", {"level": 2, "after": 2}),
    ("開発中、AI コーディング支援ツールの出力にレビュー本文が一部表示された件の扱い", {"level": 2, "after": 12}),
    ("5. 題目、提出日、発表時間", {"bold": True, "after": 6}),
], size=17)
notes(s, "要商量：停实验转写作？强底座无效怎么写？GPT-4o 是否申请？数据条款（通知、删除、AI 工具输出露出原文）？题目/截止日/发表时长。")

# speaker notes: the Japanese script of each page followed by its Chinese translation (script_content.py)
import sys  # noqa: E402
sys.path.insert(0, str(D))
from script_content import SCRIPT, TRANSLATION  # noqa: E402
assert len(SCRIPT) == len(prs.slides)
for slide, (page, _title, _sec, paras, gist) in zip(prs.slides, SCRIPT):
    ja = "\n\n".join(paras)
    zh = "\n\n".join(TRANSLATION[page])
    notes(slide, f"{ja}\n\n――――― 中文翻译 ―――――\n\n{zh}\n\n要点：{gist}")

prs.save(D / "progress_report_2026-10-15.pptx")
print("saved", D / "progress_report_2026-10-15.pptx", "| XLM-R large:", XL_DELTA)

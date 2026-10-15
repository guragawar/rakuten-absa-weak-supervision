"""Advisor meeting deck (final semester, 1st report). Numbers: review-ai/docs/thesis_numbers.md."""
import sys
sys.path.insert(0, "/tmp/claude-1000/-home-etl-gujincheng/e919a08c-08a1-4ff5-b3bf-c8ebc8a74914/scratchpad/pylib")
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

D = "/home/etl/gujincheng/meetings/2026-10-01/"
FONT = "Yu Gothic"
INK, INK2, MUTED = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x8A, 0x89, 0x83)
ACCENT, ACCENT_BG = RGBColor(0x1C, 0x5C, 0xAB), RGBColor(0xEA, 0xF2, 0xFC)
WARN, WARN_BG = RGBColor(0xB5, 0x47, 0x1C), RGBColor(0xFD, 0xF0, 0xE8)
ROW_ALT, WHITE = RGBColor(0xF5, 0xF4, 0xF1), RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def style(run, size=16, bold=False, color=INK):
    f = run.font
    f.size, f.bold, f.color.rgb, f.name = Pt(size), bold, color, FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text(slide, x, y, w, h, paras, size=16, color=INK, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, spacing=6):
    """paras: list of str | (str, dict) where dict may hold size/bold/color/level."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, p in enumerate(paras):
        s, o = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(o.get("after", spacing))
        lvl = o.get("level", 0)
        prefix = {0: "", 1: "・", 2: "　– "}[lvl] if o.get("bullet", lvl > 0) else ""
        r = para.add_run()
        r.text = prefix + s
        style(r, o.get("size", size - 2 * (lvl == 2)), o.get("bold", False), o.get("color", color))
    return tb


def box(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    return shp


def header(slide, title, num, kicker=None):
    box(slide, 0, 0, 0.18, 7.5, ACCENT)
    if kicker:
        text(slide, 0.6, 0.28, 11, 0.4, [(kicker, {"size": 13, "color": ACCENT, "bold": True})])
    text(slide, 0.6, 0.55, 12, 0.8, [(title, {"size": 28, "bold": True})])
    text(slide, 12.3, 7.0, 0.8, 0.35, [(str(num), {"size": 11, "color": MUTED})], align=PP_ALIGN.RIGHT)


def table(slide, x, y, w, rows, col_w, size=14, row_h=0.46, highlight=(), sub_small=True):
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))).table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = ACCENT
            elif i in highlight:
                c.fill.fore_color.rgb = ACCENT_BG
            else:
                c.fill.fore_color.rgb = ROW_ALT if i % 2 == 0 else WHITE
            tf = c.text_frame
            tf.word_wrap = True
            lines = val.split("\n")
            for k, ln in enumerate(lines):
                para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                r = para.add_run()
                r.text = ln
                small = sub_small and k > 0
                style(r, size - 2 if small else size, bold=(i == 0 or (i in highlight and j == 0)),
                      color=WHITE if i == 0 else (INK2 if small else INK))
    return t


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def callout(slide, x, y, w, h, title, body, warn=False):
    fg, bg = (WARN, WARN_BG) if warn else (ACCENT, ACCENT_BG)
    box(slide, x, y, w, h, bg)
    box(slide, x, y, 0.08, h, fg)
    text(slide, x + 0.25, y + 0.15, w - 0.4, h - 0.3,
         [(title, {"size": 15, "bold": True, "color": fg})] + [(b, {"size": 14}) for b in body], spacing=4)


# ---------------------------------------------------------------- 1 title
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, ACCENT)
text(s, 0.9, 1.3, 11.5, 0.5, [("修士論文 最終学期 第1回 研究報告", {"size": 18, "color": RGBColor(0xCD, 0xE2, 0xFB), "bold": True})])
text(s, 0.9, 2.0, 11.6, 2.4, [
    ("評価点を弱教師とした事前学習による", {"size": 34, "bold": True, "color": WHITE, "after": 2}),
    ("日本語ホテルレビューのアスペクト感情分析", {"size": 34, "bold": True, "color": WHITE, "after": 10}),
    ("― 暗示的なネガティブ表現の検出に向けて ―（仮題）", {"size": 20, "color": RGBColor(0xCD, 0xE2, 0xFB)}),
])
text(s, 0.9, 5.6, 11, 1.0, [
    ("2026年10月1日", {"size": 16, "color": WHITE, "after": 2}),
    ("Gu Jincheng", {"size": 16, "color": WHITE}),
])
notes(s, "开场：今天是最后一学期第一次报告。实验在暑假已经全部完成，今天主要想请老师帮我确定几件写论文前必须定下来的事。题目是暂定的，也想听老师意见。")

# ---------------------------------------------------------------- 2 purpose
s = prs.slides.add_slide(BLANK)
header(s, "本日の目的", 2)
text(s, 0.6, 1.6, 5.6, 4.8, [
    ("ご報告", {"size": 20, "bold": True, "color": ACCENT, "after": 10}),
    ("夏学期に実験はすべて完了しました", {"level": 1, "after": 10}),
    ("主な結果と、当初の仮説と異なった点", {"level": 1, "after": 10}),
    ("修論完成までの計画とリスク", {"level": 1, "after": 10}),
], size=18)
box(s, 6.6, 1.5, 6.2, 4.4, ACCENT_BG)
box(s, 6.6, 1.5, 0.08, 4.4, ACCENT)
text(s, 6.95, 1.65, 5.7, 5.0, [
    ("本日決めていただきたいこと", {"size": 20, "bold": True, "color": ACCENT, "after": 12}),
    ("① 主結果にどの実験構成を置くか", {"after": 9}),
    ("② RQ3（仮説と逆の結果）の扱い", {"after": 9}),
    ("③ データ量・年代実験の位置づけ", {"after": 9}),
    ("④ 論文の言語・体裁・中間提出日", {"after": 9}),
    ("⑤ レビュー本文の引用可否", {"after": 9}),
    ("⑥ 追加実験・対外発表の要否", {"after": 9}),
    ("→ p.14 に選択肢と私の提案をまとめています", {"size": 14, "color": INK2}),
], size=17)
notes(s, "先把今天要请老师决定的 6 件事亮出来，让老师边听边想。报告部分会讲得快一些，重点在第 10 页之后。")

# ---------------------------------------------------------------- 3 recap
s = prs.slides.add_slide(BLANK)
header(s, "研究の問いと手法（再確認）", 3)
text(s, 0.6, 1.5, 12.2, 1.6, [
    ("課題：Gold データの感情ラベルの 86.5% は暗示的（感情語を含まない）。ネガティブでは約 92%", {"level": 1}),
    ("ベースライン（E1）では、暗示的ネガティブの F1 が明示的より 4.3pp 低い", {"level": 1}),
    ("着想：楽天トラベルのレビューには投稿者による観点別の評価点がある → 約 700 万件の「無料の弱ラベル」", {"level": 1}),
], size=17)
steps = [
    ("E7 弱教師事前学習", "パネルの観点別評価点\n（低評価＝ネガ）で\n6 ヘッドを学習"),
    ("E8 Gold 微調整", "14 ラベル ABSA\n（7 観点 × 正/負）"),
    ("評価", "Gold test\nE1（事前学習なし）と比較"),
]
for i, (t_, b_) in enumerate(steps):
    x = 0.8 + i * 4.15
    box(s, x, 3.55, 3.5, 2.0, ACCENT_BG, line=ACCENT)
    text(s, x + 0.1, 3.7, 3.3, 1.8, [(t_, {"size": 18, "bold": True, "color": ACCENT, "after": 8})] + [(ln, {"size": 15, "after": 0}) for ln in b_.split("\n")], align=PP_ALIGN.CENTER)
    if i < 2:
        text(s, x + 3.5, 4.15, 0.65, 0.7, [("→", {"size": 30, "color": ACCENT, "bold": True})], align=PP_ALIGN.CENTER)
text(s, 0.8, 5.9, 12, 0.9, [
    ("データの健全性：パネルと Gold の重複 11,369 件（完全一致 9,911＋近似 1,458）を除外。Gold の約 79% がパネル由来であることを発見し、テスト汚染を防止", {"size": 14, "color": INK2}),
])
notes(s, "帮老师回忆研究设计。强调一个老师可能不知道的发现：Gold 约 79% 的评论来自面板数据，不排除会污染测试集，所以做了精确+近似去重。")

# ---------------------------------------------------------------- 4 progress
s = prs.slides.add_slide(BLANK)
header(s, "夏学期の進捗：実験はすべて完了", 4)
table(s, 0.6, 1.55, 12.1, [
    ["段階", "内容", "状態"],
    ["データ整備", "全 16 年・約 700 万件を解析、列対応を公式仕様で確定、Gold 重複を除外", "完了"],
    ["主実験", "E7 → E8、有意性検定（paired bootstrap）、複数 seed、アブレーション 4 種", "完了"],
    ["分析実験", "データ量（58 万〜676 万件）、年代（近年 / 早年）", "完了"],
    ["追加実験", "日本語専用モデル比較、積み上げ、対照学習（SCL）との比較", "完了"],
    ["未実施", "RQ4（評価点と本文の一致度）、事例分析", "執筆前に実施（約 1 週間）"],
], [2.0, 7.6, 2.5], size=15, row_h=0.62)
text(s, 0.6, 5.7, 12, 0.8, [("残りは主に執筆。GPU を要する実験は原則として不要", {"size": 16, "bold": True, "color": ACCENT})])
notes(s, "一页交代进度：实验全部完成，只剩 RQ4 和案例分析两项离线分析，约一周，不需要 GPU。")

# ---------------------------------------------------------------- 5 main result
s = prs.slides.add_slide(BLANK)
header(s, "主結果：弱教師事前学習はネガティブ検出を改善する", 5, kicker="A1  E1 vs E8（XLM-R）")
table(s, 0.6, 1.6, 12.1, [
    ["指標", "E1", "E8", "差 Δ", "95% CI", "判定"],
    ["ネガティブ集約 F1", "0.7023", "0.7223", "+2.00pp", "[+0.80, +3.21]", "有意"],
    ["マイクロ F1", "0.8234", "0.8303", "+0.70pp", "[+0.20, +1.23]", "有意"],
    ["ポジティブ集約 F1", "0.8608", "0.8637", "+0.30pp", "[−0.25, +0.83]", "非有意"],
], [3.0, 1.5, 1.5, 1.6, 3.2, 1.3], size=15, row_h=0.55, highlight=(1,))
callout(s, 0.6, 4.05, 5.9, 1.9, "頑健性", [
    "E1 全 7 回：0.7075 ± 0.0048",
    "E8 3 seed：0.7235 ± 0.0052",
    "→ 差 +1.60pp。E8 の最低値が E1 の最高値を上回る",
])
callout(s, 6.8, 4.05, 5.9, 1.9, "解釈", [
    "ネガのみ有意、ポジは非有意",
    "→ 全体の底上げではなく、ネガティブに特異的に効いている",
    "注：同一 seed でも GPU の非決定性で最大約 0.7pp 揺れるため、単一 run ではなく分布で比較",
])
notes(s, "表是 seed42 的配对 bootstrap。右下的注意很重要：我发现同一个 seed 重跑会差最多约 0.7pp（XLM-R 0.65pp、DeBERTa 0.71pp），所以改用全部 7 次运行作基线分布，结论仍然成立（E8 最差也比 E1 最好高）。主动说出来比被问到好。")

# ---------------------------------------------------------------- 6 per aspect
s = prs.slides.add_slide(BLANK)
header(s, "どこに効いたか：評価項目がある観点に集中", 6, kicker="観点別ネガティブ F1（E1 → E8）")
s.shapes.add_picture(D + "chart_aspect.png", Inches(0.5), Inches(1.5), width=Inches(7.6))
text(s, 8.4, 1.8, 4.5, 5, [
    ("改善は、パネルの評価項目が直接対応する観点（立地・部屋・設備）に集中", {"level": 1, "after": 14}),
    ("朝食・夕食はほぼ改善なし：パネルは「食事」1 項目のみ → 粒度のミスマッチ", {"level": 1, "after": 14}),
    ("立地は事前学習では最も弱いヘッド（F1 0.35）だったが、下流では最大の改善", {"level": 1, "after": 6}),
    ("→ 事前学習タスクの精度 ≠ 下流での価値", {"level": 0, "color": ACCENT, "bold": True}),
], size=16)
notes(s, "这页是机制证据：提升集中在面板有对应评分项的方面，朝食/夕食因为面板只有一个“食事”项所以没改善。这说明提升确实来自弱标签信号，而不是随机波动。")

# ---------------------------------------------------------------- 7 ablations
s = prs.slides.add_slide(BLANK)
header(s, "アブレーション：主構成の各選択は妥当", 7)
table(s, 0.6, 1.55, 12.1, [
    ["比較", "条件", "ネガ F1", "マイクロ F1", "結論"],
    ["A2 微調整方法", "全パラメータ（主構成）", "0.7224", "0.8303", "最良"],
    ["", "層別学習率", "0.7031", "0.8178", "E1（0.7033）と同等 → 利点が消える"],
    ["", "エンコーダ凍結", "0.6040", "0.6277", "下限"],
    ["A3 事前学習 epoch", "1 epoch", "0.7194", "0.8317", "2 epoch とほぼ同じ → 1 epoch で十分"],
    ["A4 総合評価ヘッド", "7 ヘッド（総合を追加）", "0.7215", "0.8300", "効果なし → 6 ヘッドで正しい"],
], [2.4, 3.1, 1.3, 1.5, 3.8], size=15, row_h=0.6, highlight=(1,))
text(s, 0.6, 5.5, 12, 1.2, [
    ("発見：層別学習率では弱教師の利点が消える → 事前学習の表現は「エンコーダを十分に更新できて初めて活きる」", {"level": 1}),
], size=16)
notes(s, "消融都是 seed42 单次。重点讲分层学习率那一行：弱监督的好处完全消失，这是一个有意思的机制观察。")

# ---------------------------------------------------------------- 8 scale & year
s = prs.slides.add_slide(BLANK)
header(s, "データ量と年代の影響（分析的実験）", 8)
s.shapes.add_picture(D + "chart_scale.png", Inches(0.5), Inches(1.5), width=Inches(7.6))
text(s, 8.4, 1.7, 4.5, 3.0, [
    ("量：200 万件で最大、676 万件ではかえって低下", {"level": 1, "after": 12}),
    ("年代：同規模でも 近年 +3.12pp vs 早年 +1.55pp", {"level": 1, "after": 6}),
    ("テスト（2018–19 年）に文体が近いデータほど有効", {"level": 2, "after": 6}),
    ("評価点ラベルの分布は年代間で大差なし → ラベルではなく言語ドメインの差", {"level": 2}),
], size=16)
callout(s, 8.4, 4.9, 4.5, 1.9, "注意（ご相談③）", [
    "各点は単一 run かつテストで比較 → 最適構成の「選択」には使わず、分析として報告したい",
], warn=True)
notes(s, "要主动说明：这些点都是单次运行、而且是在 test 上比较的；valid 指标没保存、checkpoint 已删，无法事后在 valid 上重选。所以我建议把它定位为分析性实验，不用来选主配置。另外年份效应的机制我排除过“标签分布漂移”：各年的负面率其实差不多。")

# ---------------------------------------------------------------- 9 backbone & stacking
s = prs.slides.add_slide(BLANK)
header(s, "ベースモデル比較と積み上げ", 9)
table(s, 0.6, 1.55, 7.3, [
    ["構成", "ネガ集約 F1", "E1 比"],
    ["XLM-R  E1（5 seed）", "0.7066", "—"],
    ["DeBERTa-ja  E1（5 seed）", "0.7255", "+1.9pp"],
    ["DeBERTa-ja  + E7/E8（3 seed）", "0.7447", "+3.8pp"],
], [3.9, 1.7, 1.7], size=15, row_h=0.58, highlight=(4,))
text(s, 0.6, 4.1, 7.3, 1.2, [
    ("注：数値は保存した予測からの再計算（bootstrap と同じ定義）。E7/E8 は近年 200 万件で事前学習。XLM-R ＋ E7/E8 の結果は p.5", {"size": 12, "color": INK2}),
], size=12)
text(s, 8.2, 1.55, 4.7, 5.3, [
    ("日本語専用モデル（DeBERTa-ja）は暗示ネガ F1 を有意に改善（+0.98pp, P=0.99）", {"level": 1, "after": 6}),
    ("手がかりの少ない文（短文・否定形なし）で差が大きい", {"level": 2, "after": 12}),
    ("DeBERTa 上でも E7/E8 が +1.85pp（有意）→ 2 つの改善は相補的", {"level": 1, "after": 12}),
    ("否定的結果：対照学習（SCL, SCAPT 核心）は BCE より −1.87pp（有意に劣る）", {"level": 1, "after": 6}),
    ("文書単位の 1 ビット vs 観点別の 6 ビット", {"level": 2}),
], size=16)
notes(s, "审查老师一定会问：换了更强的基座模型，弱监督还有没有用？答案是有，DeBERTa 上仍然显著 +1.85pp。SCL 对比学习是一个做过的对照，结果是负面的，也放进论文。")

# ---------------------------------------------------------------- 10 RQ answers
s = prs.slides.add_slide(BLANK)
header(s, "研究課題（RQ）ごとの回答", 10)
table(s, 0.6, 1.5, 12.1, [
    ["RQ", "問い", "回答"],
    ["RQ1", "暗示的な観点感情は明示的より難しいか", "Yes（−4.3pp）。ただし差の指標自体は seed 間で揺れが大きい"],
    ["RQ2", "評価点による弱教師事前学習は ABSA を改善するか", "Yes（+1.60pp、有意・頑健）"],
    ["RQ3", "改善は暗示的ネガティブで特に大きいか", "No（当初仮説と逆）。暗示ネガを有意に改善したのは日本語専用モデル → ご相談②"],
    ["RQ4", "評価点と本文はどの程度一致するか（弱ラベルのノイズ）", "未分析 → 執筆前に実施"],
    ["RQ5", "oseti による明示 / 暗示の分割は信頼できるか", "人手検証は実施せず、限界として記載（決定済み）"],
    ["RQ6", "LLM は暗示的な表現で優れるか", "実施せず（範囲外）"],
], [1.0, 5.0, 6.1], size=14, row_h=0.64, highlight=(3,))
notes(s, "RQ3 是今天最需要老师判断的点：原假设是弱监督对隐式负面特别有效，结果不成立。我的想法是把它当作正式的负面发现写，并和“日语专用模型确实改善隐式负面”对比，这样故事反而更完整。")

# ---------------------------------------------------------------- 11 thesis outline
s = prs.slides.add_slide(BLANK)
header(s, "論文の構成案", 11)
table(s, 0.6, 1.5, 12.1, [
    ["章", "内容", "素材の状況"],
    ["1  序論", "動機、RQ", "あり"],
    ["2  関連研究", "ABSA、弱教師、ドメイン適応事前学習など 27 本", "整理済み → 論文調に書き直し"],
    ["3  データ", "Gold、パネル、重複除外、列対応の確定", "あり（図は未作成）"],
    ["4  手法", "2 段階学習（E7 → E8）", "あり"],
    ["5  実験", "主結果、アブレーション", "あり（図表は未作成）"],
    ["6  分析", "観点別、データ量・年代、ベースモデル、機構分析、RQ4、事例", "RQ4・事例が未実施"],
    ["7  結論", "RQ への回答、限界、今後の課題", "限界のリストは確定"],
], [2.2, 6.1, 3.8], size=15, row_h=0.6)
notes(s, "按这个章节结构写。想问老师：扩展实验（骨干模型、SCL）放第 6 章分析，还是单独成章。")

# ---------------------------------------------------------------- 12 risks
s = prs.slides.add_slide(BLANK)
header(s, "リスクと対応", 12)
table(s, 0.6, 1.5, 12.1, [
    ["リスク", "対応"],
    ["テストセットで構成を比較した（データ量・年代）", "分析的実験として位置づけ、主表は事前に決めた構成を使う"],
    ["GPU の非決定性（同一 seed で最大約 0.7pp）", "全 run の分布で比較し、単一 run の差は主張しない"],
    ["oseti による明示 / 暗示分割は人手で未検証", "限界に明記（検証用サンプル 400 件と採点ツールは準備済み）"],
    ["2016 年のデータが不完全（約 20% のみ）", "データの章に明記"],
    ["ライセンスデータの本文引用", "要確認（ご相談⑤）"],
    ["一部の分析スクリプトが未整理", "執筆前に整理し、保存済みの予測から再計算して照合"],
], [5.4, 6.7], size=15, row_h=0.62)
notes(s, "风险都已有对策。前两条是我自己查出来的，主动讲，体现结论是经得起追问的。")

# ---------------------------------------------------------------- 13 schedule
s = prs.slides.add_slide(BLANK)
header(s, "スケジュール案（提出日から逆算）", 13)
table(s, 0.6, 1.5, 12.1, [
    ["時期", "作業", "区切り"],
    ["10 月前半（2 週）", "RQ4・事例分析、図表作成、スクリプト整理", "執筆準備の完了"],
    ["10 月後半〜11 月", "3〜5 章（データ・手法・実験）を執筆", "11 月末：中間稿を提出"],
    ["12 月", "6 章（分析）、1・2・7 章を執筆", "12 月末：全体の初稿"],
    ["1 月〜", "修正、要旨、発表準備", "提出日：要確認"],
], [2.8, 6.1, 3.2], size=15, row_h=0.66)
text(s, 0.6, 5.2, 12, 1.2, [
    ("提出日・中間提出の有無が分かり次第、この表を確定します（ご相談④）", {"level": 1}),
    ("対外発表（言語処理学会 年次大会など）をする場合は、年明けの投稿締切に合わせて 5 章を前倒し", {"level": 1}),
], size=16)
notes(s, "日期是按常规修论节奏假设的，要请老师告诉我正式提交日和中期提交节点，会后更新。NLP 年次大会的截止日期我会后再查确认。")

# ---------------------------------------------------------------- 14 decisions
s = prs.slides.add_slide(BLANK)
header(s, "本日ご相談したいこと", 14)
table(s, 0.6, 1.4, 12.1, [
    ["論点", "選択肢", "私の提案"],
    ["① 主結果の構成", "A：XLM-R の主構成を主表、DeBERTa の積み上げは拡張実験\nB：DeBERTa ＋ E7/E8 を主表", "A（事前に決めた設定で、最も頑健）"],
    ["② RQ3 の扱い", "A：否定的な発見として正式に回答\nB：分析の章での議論にとどめる", "A（ベースモデルの結果と対比）"],
    ["③ データ量・年代", "A：分析的実験として報告\nB：主張に含める（valid で選び直す追加実験が必要）", "A"],
    ["④ 言語・体裁", "日本語 / 英語、テンプレート、分量、中間提出日", "ご指示ください"],
    ["⑤ 本文の引用", "原文を引用 / 言い換えた例のみ", "協定の確認方法をご相談"],
    ["⑥ 追加実験・発表", "追加実験なしで執筆に集中 / 学会発表", "追加実験なし。発表はご判断"],
], [2.3, 6.3, 3.5], size=14, row_h=0.72, sub_small=False)
notes(s, "逐项问，老师回答后当场在自己的纸上记下结论。①③是连着的：如果老师选 B，就需要在 valid 上重新选配置，要补实验，时间会紧。")

# ---------------------------------------------------------------- 15 appendix
s = prs.slides.add_slide(BLANK)
header(s, "付録：数値の出典と定義", 15)
text(s, 0.6, 1.6, 12, 5, [
    ("すべての数値は docs/thesis_numbers.md に出典（結果ファイル）とともに登録済み", {"level": 1, "after": 12}),
    ("ネガティブ集約 F1：7 つのネガティブラベルをまとめて計算した F1（ラベルごとの F1 の平均ではない）", {"level": 1, "after": 12}),
    ("有意性：paired bootstrap（2,000 回）。複数 seed の比較は seed 平均の予測で検定", {"level": 1, "after": 12}),
    ("E1 の 0.7033 と 0.7023 の差：学習時の評価と保存した予測からの再計算で、バッチ処理が異なるため（≤ 0.001）", {"level": 1, "after": 12}),
    ("評価には Gold test（5,356 文）のみを使用。パネルとの重複は除外済み", {"level": 1}),
], size=16)
notes(s, "备用页，被问到数字口径时翻到这里。")

out = D + "advisor_meeting_2026-10-01.pptx"
prs.save(out)
print(out)
